#!/usr/bin/env python3
"""
Test script to debug Excel logging issues
"""

import os
import sys
from datetime import datetime

# Add the current directory to Python path
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

try:
    from sm import AIEventLogger, AIMonitorConfig, CMDStyler
except ImportError as e:
    print(f"Import error: {e}")
    sys.exit(1)

def test_excel_logging():
    """Test Excel logging functionality"""
    print("🧪 Testing Excel Logging Functionality")
    print("=" * 50)
    
    try:
        # Create config
        config = AIMonitorConfig()
        print(f"✓ Config created")
        print(f"  Log file: {config.log_file}")
        
        # Create logger
        logger = AIEventLogger(config)
        print(f"✓ Logger created")
        
        # Test basic event logging
        print("\n📝 Testing basic event logging...")
        
        test_events = [
            ("File Access", "Test file access detected", "test.exe", "C:\\test\\file.txt", "Low"),
            ("AI Tool Usage", "ChatGPT usage detected", "chrome.exe", "https://chat.openai.com", "Medium"),
            ("Process Activity", "AI process started", "python.exe", "C:\\ai\\script.py", "High"),
            ("Network Connection", "AI service connection", "msedge.exe", "https://claude.ai", "Critical")
        ]
        
        for i, (event_type, description, process, file_url, risk) in enumerate(test_events):
            print(f"  Logging event {i+1}: {event_type}")
            logger.log_event(event_type, description, process, file_url, risk)
        
        # Force save
        print("\n💾 Forcing Excel save...")
        logger.save_workbook()
        
        # Check if file exists
        if os.path.exists(config.log_file):
            print(f"✓ Excel file created: {config.log_file}")
            file_size = os.path.getsize(config.log_file)
            print(f"  File size: {file_size} bytes")
        else:
            print(f"✗ Excel file NOT found: {config.log_file}")
            return False
        
        # Test summary sheet creation
        print("\n📊 Testing summary sheet creation...")
        logger.create_summary_sheet()
        
        print("\n✅ Excel logging test completed successfully!")
        return True
        
    except Exception as e:
        print(f"✗ Error during testing: {e}")
        import traceback
        traceback.print_exc()
        return False

def test_excel_file_content():
    """Test reading Excel file content"""
    print("\n📖 Testing Excel file content...")
    
    try:
        from openpyxl import load_workbook
        
        config = AIMonitorConfig()
        
        if not os.path.exists(config.log_file):
            print(f"✗ Excel file not found: {config.log_file}")
            return False
        
        # Load workbook
        wb = load_workbook(config.log_file)
        print(f"✓ Workbook loaded successfully")
        
        # List sheets
        print(f"  Sheets: {wb.sheetnames}")
        
        # Check main sheet
        if "AI Activity Monitor" in wb.sheetnames:
            sheet = wb["AI Activity Monitor"]
            print(f"✓ Main sheet found")
            
            # Count rows
            row_count = sheet.max_row
            print(f"  Total rows: {row_count}")
            
            if row_count > 1:
                print("  Sample data:")
                for row in range(1, min(6, row_count + 1)):
                    row_data = []
                    for col in range(1, 7):  # First 6 columns
                        cell_value = sheet.cell(row=row, column=col).value
                        row_data.append(str(cell_value) if cell_value else "")
                    print(f"    Row {row}: {' | '.join(row_data)}")
            else:
                print("  ⚠️  No data rows found (only headers)")
        
        # Check summary sheet
        if "AI_Summary" in wb.sheetnames:
            sheet = wb["AI_Summary"]
            print(f"✓ Summary sheet found")
            row_count = sheet.max_row
            print(f"  Summary rows: {row_count}")
        
        return True
        
    except Exception as e:
        print(f"✗ Error reading Excel file: {e}")
        import traceback
        traceback.print_exc()
        return False

if __name__ == "__main__":
    print("🔍 Excel Logging Debug Tool")
    print("=" * 50)
    
    # Test logging
    if test_excel_logging():
        # Test content
        test_excel_file_content()
        
        print("\n🎉 All tests completed!")
        print(f"📁 Check the Excel file: {AIMonitorConfig().log_file}")
    else:
        print("\n❌ Tests failed!")
        sys.exit(1)
